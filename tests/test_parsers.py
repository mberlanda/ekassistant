import pathlib

import docx
import pytest

from ekassistant.ingest.chunker import chunk_document
from ekassistant.ingest.parsers import parse_document
from ekassistant.ingest.parsers._shared import escape_accidental_heading
from ekassistant.ingest.parsers.docx_parser import parse as parse_docx
from ekassistant.ingest.parsers.html_parser import parse as parse_html
from ekassistant.ingest.parsers.pdf_parser import _escape_lines
from ekassistant.ingest.parsers.pdf_parser import parse as parse_pdf
from ekassistant.ingest.parsers.text_parser import parse as parse_text

_SEED_CORPUS = pathlib.Path(__file__).parent.parent / "seed_corpus"


def test_text_parser_returns_file_contents_unchanged(tmp_path):
    path = tmp_path / "doc.md"
    path.write_text("# Heading\n\nBody text.")

    assert parse_text(path) == "# Heading\n\nBody text."


def test_html_parser_converts_headings_and_paragraphs(tmp_path):
    path = tmp_path / "doc.html"
    path.write_text(
        "<html><body>"
        "<h1>Title</h1>"
        "<h2>Section One</h2>"
        "<p>First paragraph.</p>"
        "<h2>Section Two</h2>"
        "<p>Second paragraph.</p>"
        "</body></html>"
    )

    text = parse_html(path)

    assert "# Title" in text
    assert "## Section One" in text
    assert "First paragraph." in text
    assert "## Section Two" in text
    assert "Second paragraph." in text
    # Document order is preserved, not grouped by tag type.
    assert text.index("Section One") < text.index("Section Two")


def test_html_parser_ignores_non_heading_non_paragraph_tags(tmp_path):
    path = tmp_path / "doc.html"
    path.write_text(
        "<html><body><nav>Skip this nav</nav><h1>Title</h1>"
        "<p>Keep this.</p><script>console.log('skip')</script></body></html>"
    )

    text = parse_html(path)

    assert "Skip this nav" not in text
    assert "console.log" not in text
    assert "Keep this." in text


def test_html_parser_keeps_a_space_around_nested_inline_tags(tmp_path):
    path = tmp_path / "doc.html"
    path.write_text(
        '<html><body><p>Contact <a href="x">the on-call engineer</a>immediately.</p>'
        "</body></html>"
    )

    text = parse_html(path)

    assert text == "Contact the on-call engineer immediately."


def test_html_parser_escapes_paragraph_text_that_looks_like_a_heading(tmp_path):
    # A <p> is never a real heading, but "# 1 priority item" as its text
    # would otherwise be misread as one by chunker.py's _HEADING_RE.
    path = tmp_path / "doc.html"
    path.write_text("<html><body><p># 1 priority item</p></body></html>")

    text = parse_html(path)

    assert text == "\\# 1 priority item"
    # And, end to end: the chunker must not treat it as a heading boundary.
    chunks = chunk_document("doc.html", "doc.html", text, [])
    assert len(chunks) == 1
    assert chunks[0].text == "\\# 1 priority item"


def test_html_parser_skips_empty_tags(tmp_path):
    path = tmp_path / "doc.html"
    path.write_text("<html><body><h2>   </h2><p>Real content.</p></body></html>")

    text = parse_html(path)

    assert text.strip() == "Real content."


def test_docx_parser_maps_heading_styles_to_hash_prefixes(tmp_path):
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_heading("Subsection", level=2)
    document.add_paragraph("Body text under the subsection.")
    document.save(path)

    text = parse_docx(path)

    assert "# Title" in text
    assert "## Subsection" in text
    assert "Body text under the subsection." in text


def test_docx_parser_skips_empty_paragraphs(tmp_path):
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("")
    document.add_paragraph("Real content.")
    document.save(path)

    text = parse_docx(path)

    assert text == "# Title\n\nReal content."


def test_docx_parser_escapes_normal_paragraph_text_that_looks_like_a_heading(tmp_path):
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("# 1234: ticket reference, not a heading")
    document.save(path)

    text = parse_docx(path)

    assert "\\# 1234: ticket reference, not a heading" in text
    # End to end: only the real "Title" heading should create a section
    # boundary - the ticket-reference paragraph must not create a second.
    chunks = chunk_document("doc.docx", "doc.docx", text, [])
    assert len(chunks) == 1
    assert "1234" in chunks[0].text


def test_docx_parser_escapes_heading_like_text_on_a_manual_line_break(tmp_path):
    # python-docx renders a manual line break (Shift+Enter) within one
    # paragraph as an embedded "\n" in .text - a later line can start
    # with "# " even when the paragraph itself doesn't.
    path = tmp_path / "doc.docx"
    document = docx.Document()
    paragraph = document.add_paragraph("First line")
    paragraph.add_run().add_break()
    paragraph.add_run("# Second line looks like a heading")
    document.save(path)

    text = parse_docx(path)

    assert "First line" in text
    assert "\\# Second line looks like a heading" in text
    # End to end: the chunker must not split here or treat it as a heading.
    chunks = chunk_document("doc.docx", "doc.docx", text, [])
    assert len(chunks) == 1


def test_escape_accidental_heading_leaves_normal_text_untouched():
    assert escape_accidental_heading("Ordinary sentence.") == "Ordinary sentence."
    assert escape_accidental_heading("#hashtag-no-space") == "#hashtag-no-space"


def test_escape_accidental_heading_escapes_all_heading_levels():
    for level in range(1, 7):
        text = f"{'#' * level} looks like a heading"
        assert escape_accidental_heading(text) == f"\\{text}"


def test_pdf_escape_lines_escapes_only_lines_that_look_like_headings():
    text = "Normal line.\n# 1 looks like a heading\nAnother normal line."

    escaped = _escape_lines(text)

    assert escaped == "Normal line.\n\\# 1 looks like a heading\nAnother normal line."


def test_pdf_parser_extracts_text_from_the_real_seed_corpus_fixture():
    # No lightweight, dependency-free way to author a synthetic PDF fixture
    # in a unit test (unlike .md/.html/.docx, which are just text/XML this
    # test can write directly) - see docs/decisions/0009 for how the
    # committed seed_corpus/vendor-contract-terms.pdf was produced.
    # Testing against the real committed fixture keeps this test honest
    # about what pypdf actually extracts from a real PDF.
    text = parse_pdf(_SEED_CORPUS / "vendor-contract-terms.pdf")

    assert "Vendor Contract Terms" in text
    assert "Payment terms" in text
    assert "net-45" in text


def test_parse_document_dispatches_by_extension(tmp_path):
    md_path = tmp_path / "doc.md"
    md_path.write_text("plain markdown")

    assert parse_document(md_path) == "plain markdown"


def test_parse_document_dispatches_html_pdf_and_docx_not_just_the_raw_parse_functions(tmp_path):
    # The dispatch table itself (not just each format's parse()) is what
    # actually wires every new parser into FilesystemConnector - worth
    # exercising directly, not only indirectly through .md.
    html_path = tmp_path / "doc.html"
    html_path.write_text("<html><body><p>hello</p></body></html>")
    assert parse_document(html_path) == "hello"

    htm_path = tmp_path / "doc.htm"
    htm_path.write_text("<html><body><p>hello</p></body></html>")
    assert parse_document(htm_path) == "hello"

    docx_path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("hello")
    document.save(docx_path)
    assert parse_document(docx_path) == "hello"

    assert "Vendor Contract Terms" in parse_document(_SEED_CORPUS / "vendor-contract-terms.pdf")


def test_parse_document_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "doc.exe"
    path.write_text("irrelevant")

    with pytest.raises(ValueError, match="unsupported file extension"):
        parse_document(path)


def test_parse_document_extension_matching_is_case_insensitive(tmp_path):
    path = tmp_path / "doc.MD"
    path.write_text("plain markdown")

    assert parse_document(path) == "plain markdown"
